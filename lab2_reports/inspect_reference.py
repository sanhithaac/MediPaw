import pptx

prs = pptx.Presentation('lab2_reports/MediPaw_Team11.pptx')
print(f"Slide Dimensions: {prs.slide_width / 914400:.2f} x {prs.slide_height / 914400:.2f} inches")
print(f"Number of slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    print(f"\n============================== SLIDE {idx+1} ==============================")
    for s_idx, shape in enumerate(slide.shapes):
        pos_str = f"pos=({shape.left/914400:.2f}, {shape.top/914400:.2f}), size=({shape.width/914400:.2f}x{shape.height/914400:.2f})"
        
        # Check text
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                p_text = p.text.strip()
                if not p_text:
                    continue
                font_name = p.font.name if p.font else None
                font_size = p.font.size.pt if (p.font and p.font.size) else None
                bold = p.font.bold if p.font else None
                color_hex = None
                try:
                    if p.font and p.font.color and p.font.color.rgb:
                        color_hex = str(p.font.color.rgb)
                except Exception:
                    pass
                print(f"  [Text]: '{p_text}' | Font: {font_name}, Size: {font_size}, Bold: {bold}, Color: {color_hex} | {pos_str}")
                
        elif shape.has_table:
            t = shape.table
            print(f"  [Table]: {len(t.rows)} rows x {len(t.columns)} cols | {pos_str}")
            for r_i, row in enumerate(t.rows):
                row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                print(f"    Row {r_i}: {row_vals}")
                
        elif shape.shape_type == 13: # picture
            print(f"  [Picture]: {shape.name} | {pos_str}")
            
        else:
            fill_desc = "None"
            if hasattr(shape, 'fill'):
                try:
                    if shape.fill.type == 1:
                        fill_desc = f"Solid({shape.fill.fore_color.rgb})"
                except Exception:
                    pass
            print(f"  [Shape]: {shape.name} (type {shape.shape_type}) | fill={fill_desc} | {pos_str}")
