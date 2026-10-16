# export_review2_all.py
"""
Master orchestrator to build both individual Review 2 presentations and export matching PDFs:
1. MediPaw_Review2_AC_Sanhitha_Reddy.pptx (44 slides)
2. MediPaw_Review2_AC_Sanhitha_Reddy.pdf
3. MediPaw_Review2_Shaik_Abdus_Sattar.pptx (44 slides)
4. MediPaw_Review2_Shaik_Abdus_Sattar.pdf
"""

import os
import sys
from pptx import Presentation
from build_sanhitha_deck import build_sanhitha_deck
from build_sattar_deck import build_sattar_deck

def export_pptx_to_pdf(pptx_abs_path, pdf_abs_path):
    print(f"[PDF Export] Converting {os.path.basename(pptx_abs_path)} -> {os.path.basename(pdf_abs_path)}...")
    try:
        import win32com.client
        powerpoint = win32com.client.Dispatch("PowerPoint.Application")
        # Open presentation without displaying window
        deck = powerpoint.Presentations.Open(pptx_abs_path, WithWindow=False)
        # 32 corresponds to ppSaveAsPDF
        deck.SaveAs(pdf_abs_path, 32)
        deck.Close()
        powerpoint.Quit()
        if os.path.exists(pdf_abs_path) and os.path.getsize(pdf_abs_path) > 0:
            print(f"[SUCCESS] Exported PDF ({os.path.getsize(pdf_abs_path):,} bytes): {pdf_abs_path}")
            return True
        else:
            print(f"[ERROR] PDF file not found or empty after export: {pdf_abs_path}")
            return False
    except Exception as e:
        print(f"[ERROR] PowerPoint COM export failed: {e}")
        return False

def main():
    base_dir = os.path.abspath("lab2_reports")
    os.makedirs(base_dir, exist_ok=True)

    sanhitha_pptx = os.path.join(base_dir, "MediPaw_Review2_AC_Sanhitha_Reddy.pptx")
    sanhitha_pdf = os.path.join(base_dir, "MediPaw_Review2_AC_Sanhitha_Reddy.pdf")
    sattar_pptx = os.path.join(base_dir, "MediPaw_Review2_Shaik_Abdus_Sattar.pptx")
    sattar_pdf = os.path.join(base_dir, "MediPaw_Review2_Shaik_Abdus_Sattar.pdf")

    print("\n" + "="*70)
    print("STEP 1: BUILDING SANHITHA REDDY REVIEW 2 DECK")
    print("="*70)
    build_sanhitha_deck(sanhitha_pptx)
    prs1 = Presentation(sanhitha_pptx)
    print(f"Verified Sanhitha Deck: {len(prs1.slides)} slides (Size: {os.path.getsize(sanhitha_pptx):,} bytes)")

    print("\n" + "="*70)
    print("STEP 2: BUILDING SHAIK ABDUS SATTAR REVIEW 2 DECK")
    print("="*70)
    build_sattar_deck(sattar_pptx)
    prs2 = Presentation(sattar_pptx)
    print(f"Verified Sattar Deck: {len(prs2.slides)} slides (Size: {os.path.getsize(sattar_pptx):,} bytes)")

    print("\n" + "="*70)
    print("STEP 3: EXPORTING PDF COPIES VIA POWERPOINT COM")
    print("="*70)
    export_pptx_to_pdf(sanhitha_pptx, sanhitha_pdf)
    export_pptx_to_pdf(sattar_pptx, sattar_pdf)

    print("\n" + "="*70)
    print("VERIFICATION SUMMARY")
    print("="*70)
    for p in [sanhitha_pptx, sanhitha_pdf, sattar_pptx, sattar_pdf]:
        status = "EXISTS" if os.path.exists(p) else "MISSING"
        size = f"{os.path.getsize(p):,} bytes" if os.path.exists(p) else "0 bytes"
        print(f"• {os.path.basename(p)}: [{status}] ({size}) -> {p}")
    print("="*70 + "\n")

if __name__ == "__main__":
    main()
